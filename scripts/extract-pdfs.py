#!/usr/bin/env python3
"""
Baut die Vokabeldaten aus den ukrainischen Klett-PDFs (seed/*.pdf) neu auf.

Läuft in mehreren, einzeln wiederholbaren Schritten über seed/dataset.json
(alles unter seed/ ist gitignored — Vokabelinhalte dürfen nie ins Repo):

  extract   PDFs → seed/dataset.json; übernimmt Beispielsätze aus der alten
            Excel (Vokabeln_GreenLine_2021_gesamt.xlsx) und schreibt die noch
            offenen Wortpaare dedupliziert nach seed/missing-sentences.json
  merge F…  generierte Satz-JSONs ({start: Index in missing-sentences.json,
            pairs: [[en, de], …]}) in seed/dataset.json einarbeiten
  write     seed/dataset.json (+ optional seed/lesson-names.json) →
            Vokabeln_GreenLine_BW2016_gesamt.xlsx

Interpreter: venv mit pdfplumber + openpyxl (siehe README des Scratchpads),
Aufruf aus dem Projektroot.
"""
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SEED = ROOT / "seed"
DATASET = SEED / "dataset.json"
MISSING = SEED / "missing-sentences.json"
LESSON_NAMES = SEED / "lesson-names.json"
OLD_XLSX = ROOT / "Vokabeln_GreenLine_2021_gesamt.xlsx"
NEW_XLSX = ROOT / "Vokabeln_GreenLine_BW2016_gesamt.xlsx"

# (Band, Datei, erwartete Datenzeilen nach Header-Filter)
PDFS = [
    ("Band 1", "Vokabelliste_834210_GL_2014_1_ukr.116065.pdf", 904),
    ("Band 2", "Vokabelliste_834120_GL_BW_16_2_ukr.pdf", 784),
    ("Band 3", "Vokabelliste_834130_GL_BW_16_3_ukr.pdf", 741),
    ("Band 4", "Vokabelliste_834240_GL_2014_4_ukr.116095.pdf", 746),
    ("Band 5", "Vokabelliste_834250_GL_2014_5_ukr.116105.pdf", 677),
]

# ---------------------------------------------------------------- Phonetik
# Klett-Phonetik-Font → IPA. Digraphen zuerst (Max-Munch).
DIGRAPHS = {
    "A:": "ɑː", "C:": "ɔː", "3:": "ɜː", "i:": "iː", "u:": "uː",
    "aI": "aɪ", "aU": "aʊ", "eI": "eɪ", "CI": "ɔɪ", "EU": "əʊ",
    "IE": "ɪə", "eE": "eə", "UE": "ʊə", "tS": "tʃ", "dZ": "dʒ",
    # Quell-Varianten/Sonderfälle: words→wɜːdz, mainstream→striːm,
    # genre→ˈʒɑːnrə (Longman), vereinzelt kleingeschriebenes aI
    "E:": "ɜː", "I:": "iː", "R:": "ɑː", "ai": "aɪ",
}
SINGLES = {
    "!": "[", "?": "]", "*": "ˈ", "+": "ˌ", "9": "ˌ",
    "I": "ɪ", "E": "ə", "U": "ʊ", "O": "ɒ", "C": "ɔ", "A": "ɑ",
    "0": "ʌ", "x": "æ", "S": "ʃ", "Z": "ʒ", "T": "θ", "D": "ð",
    "N": "ŋ", "G": "ɡ", "g": "ɡ", "X": "x", "a": "ʌ",
}
# Zeichen, die unverändert gültiges IPA/Beiwerk sind
IDENTITY = set("bdefhijklmnprstuvwz iue'’,;.()/- …:")

def to_ipa(raw):
    out, unknown, i = [], [], 0
    while i < len(raw):
        two = raw[i:i + 2]
        if two in DIGRAPHS:
            out.append(DIGRAPHS[two])
            i += 2
            continue
        ch = raw[i]
        if ch in SINGLES:
            out.append(SINGLES[ch])
        elif ch in IDENTITY:
            out.append(ch)
        else:
            unknown.append(ch)
            out.append(ch)
        i += 1
    return "".join(out), unknown

# ---------------------------------------------------------------- Extraktion
def clean_cell(text):
    if not text:
        return ""
    t = text.replace("\xa0", " ").replace("﻿", "")
    t = re.sub(r"-\n(?=[a-zäöüß])", "", t)  # Silbentrennung am Zellen-Umbruch
    t = re.sub(r"\s*\n\s*", " ", t)
    return re.sub(r"\s+", " ", t).strip()

def has_cyrillic(s: str) -> bool:
    return any("CYRILLIC" in unicodedata.name(c, "") for c in s)

def extract():
    import pdfplumber

    rows, unknown_chars, anomalies = [], Counter(), []
    unknown_samples = {}
    for band, fname, expected in PDFS:
        pdf = pdfplumber.open(SEED / fname)
        count = 0
        for page in pdf.pages:
            for table in page.extract_tables():
                for raw in table:
                    cells = [clean_cell(c) for c in raw]
                    if len(cells) != 6 or not any(cells):
                        continue
                    if cells[0].startswith(("Lektion", "Vokabular")):
                        continue
                    code, part, english, phonetic, german, _ukr = cells
                    count += 1
                    lesson = f"{code} {part}".strip()
                    ipa, unk = to_ipa(phonetic)
                    for ch in unk:
                        unknown_chars[ch] += 1
                        unknown_samples.setdefault(ch, f"{english}: {phonetic}")
                    if not english or not german:
                        anomalies.append((band, lesson, english, german))
                    if has_cyrillic(english + german + phonetic):
                        anomalies.append((band, lesson, english, "KYRILLISCH!"))
                    rows.append({
                        "band": band, "lesson": lesson, "english": english,
                        "phonetic": ipa, "german": german,
                        "example_en": "", "example_de": "",
                    })
        assert count == expected, f"{band}: {count} Zeilen, erwartet {expected}"
        print(f"{band}: {count} Zeilen ✓ ({fname})")

    reused = reuse_sentences(rows)
    DATASET.write_text(json.dumps({"rows": rows}, ensure_ascii=False, indent=1))

    missing, seen = [], set()
    for r in rows:
        if r["example_en"]:
            continue
        key = sentence_key(r["english"], r["german"])
        if key in seen:
            continue
        seen.add(key)
        missing.append({"key": key, "band": r["band"], "lesson": r["lesson"],
                        "english": r["english"], "german": r["german"]})
    MISSING.write_text(json.dumps(missing, ensure_ascii=False, indent=1))

    print(f"\nGesamt: {len(rows)} Zeilen → {DATASET.name}")
    print(f"Beispielsätze übernommen: {reused} | offen: "
          f"{sum(1 for r in rows if not r['example_en'])} Zeilen "
          f"({len(missing)} eindeutige Wortpaare → {MISSING.name})")
    if unknown_chars:
        print("\nUnbekannte Phonetik-Zeichen:")
        for ch, n in unknown_chars.most_common():
            print(f"  {ch!r} ×{n}  z. B. {unknown_samples[ch]}")
    if anomalies:
        print(f"\nAnomalien ({len(anomalies)}):")
        for a in anomalies[:20]:
            print("  ", a)

# ------------------------------------------------------------- Beispielsätze
STOPWORDS = {"sich", "etwas", "etw", "jemand", "jemanden", "jemandem", "jdn",
             "jdm", "der", "die", "das", "den", "dem", "ein", "eine", "einen",
             "sein", "haben", "werden", "und", "oder", "mit", "von", "auf",
             "für", "aus", "nicht", "man"}

def norm(s: str) -> str:
    s = s.lower().strip().replace("’", "'").replace("‘", "'").replace("…", "...")
    return re.sub(r"\s+", " ", s)

def de_tokens(s: str) -> set[str]:
    return {t for t in re.findall(r"[a-zäöüß]+", norm(s)) if t not in STOPWORDS}

def sentence_key(en: str, de: str) -> str:
    return f"{norm(en)}||{norm(de)}"

def reuse_sentences(rows):
    from openpyxl import load_workbook

    exact = {}
    by_en = {}
    ws = load_workbook(OLD_XLSX, read_only=True)["Vokabeln"]
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        _band, _lesson, en, _ph, de, ex_en, ex_de = (str(c or "").strip() for c in row[:7])
        if not en or not ex_en or not ex_de:
            continue
        exact.setdefault((norm(en), norm(de)), (ex_en, ex_de))
        by_en.setdefault(norm(en), []).append((de, ex_en, ex_de))

    reused = 0
    for r in rows:
        pair = exact.get((norm(r["english"]), norm(r["german"])))
        if not pair:
            candidates = by_en.get(norm(r["english"]), [])
            unique = {(ex_en, ex_de) for _de, ex_en, ex_de in candidates}
            if len(unique) == 1:
                old_tokens = set().union(*(de_tokens(de) for de, _, _ in candidates))
                if old_tokens & de_tokens(r["german"]):
                    pair = unique.pop()
        if pair:
            r["example_en"], r["example_de"] = pair
            reused += 1
    return reused

def merge(files):
    data = json.loads(DATASET.read_text())
    missing = json.loads(MISSING.read_text())
    sentences = {}
    for f in files:
        blob = json.loads(Path(f).read_text())
        for j, (ex_en, ex_de) in enumerate(blob["pairs"]):
            sentences[missing[blob["start"] + j]["key"]] = (ex_en, ex_de)
    filled = 0
    for r in data["rows"]:
        if not r["example_en"]:
            pair = sentences.get(sentence_key(r["english"], r["german"]))
            if pair:
                r["example_en"], r["example_de"] = pair
                filled += 1
    DATASET.write_text(json.dumps(data, ensure_ascii=False, indent=1))
    open_rows = sum(1 for r in data["rows"] if not r["example_en"])
    print(f"{filled} Zeilen gefüllt, noch offen: {open_rows}")

# ------------------------------------------------------------------- Excel
def write():
    from openpyxl import Workbook

    data = json.loads(DATASET.read_text())
    wb = Workbook()
    ws = wb.active
    ws.title = "Vokabeln"
    ws.append(["Band", "Lektion", "Englisch", "Phonetik", "Deutsch",
               "Beispielsatz (EN)", "Beispielsatz (DE)"])
    for r in data["rows"]:
        ws.append([r["band"], r["lesson"], r["english"], r["phonetic"],
                   r["german"], r["example_en"], r["example_de"]])
    for col, width in zip("ABCDEFG", (10, 12, 30, 24, 40, 45, 45)):
        ws.column_dimensions[col].width = width

    if LESSON_NAMES.exists():
        names = json.loads(LESSON_NAMES.read_text())
        ls = wb.create_sheet("Lektionen")
        ls.append(["Lektion", "Name", "Beschreibung"])
        for item in names:
            ls.append([item["code"], item["name"], item.get("beschreibung", "")])
        for col, width in zip("ABC", (12, 40, 60)):
            ls.column_dimensions[col].width = width

    wb.save(NEW_XLSX)
    open_rows = sum(1 for r in data["rows"] if not r["example_en"])
    print(f"{NEW_XLSX.name}: {len(data['rows'])} Zeilen"
          + (f" — WARNUNG: {open_rows} ohne Beispielsatz" if open_rows else ""))

if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "extract":
        extract()
    elif cmd == "merge":
        merge(sys.argv[2:])
    elif cmd == "write":
        write()
    else:
        sys.exit("Aufruf: extract | merge <json…> | write")
